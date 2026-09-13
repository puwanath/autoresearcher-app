"""PowerPoint export (PRD §5.4): one slide per report section, native (editable) charts and tables."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from app.agents.stats import chart_brands
from app.schemas.research import AnalysisReport, ExtractedProduct, ImageAsset, PriceStats, ResearchRequest

NAVY, INK, GREY, ACCENT, WHITE = (
    RGBColor(0x0B, 0x3D, 0x5C),
    RGBColor(0x1D, 0x1D, 0x1F),
    RGBColor(0x6E, 0x6E, 0x73),
    RGBColor(0x00, 0x71, 0xE3),
    RGBColor(0xFF, 0xFF, 0xFF),
)
FONT = "Tahoma"  # ships with Windows and macOS and covers Thai
W, H = Inches(13.333), Inches(7.5)


class Deck:
    def __init__(self, title: str):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.title = title
        self.n = 0

    # ---- primitives ----------------------------------------------------
    def _text(
        self, slide, x, y, w, h, text, *, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP
    ):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        lines = text if isinstance(text, list) else [text]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = line
            r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
            r.font.color.rgb = color
        return tb

    def _bullets(self, slide, x, y, w, h, items, *, size=14, color=INK):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(6)
            r = p.add_run()
            r.text = f"•  {it}"
            r.font.size, r.font.name = Pt(size), FONT
            r.font.color.rgb = color
        return tb

    def slide(self, heading: str, eyebrow: str = ""):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        bar = s.shapes.add_shape(1, 0, 0, W, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        if eyebrow:
            self._text(s, Inches(0.6), Inches(0.35), Inches(8), Inches(0.3), eyebrow.upper(), size=10, color=GREY)
        self._text(s, Inches(0.6), Inches(0.6), Inches(12), Inches(0.7), heading, size=26, bold=True, color=NAVY)
        self._text(
            s,
            Inches(0.6),
            H - Inches(0.45),
            Inches(9),
            Inches(0.3),
            f"AutoResearch Agent — Puwanath Baibua · {self.title}",
            size=9,
            color=GREY,
        )
        self._text(
            s,
            W - Inches(1.6),
            H - Inches(0.45),
            Inches(1),
            Inches(0.3),
            str(self.n),
            size=9,
            color=GREY,
            align=PP_ALIGN.RIGHT,
        )
        return s

    def table(
        self, slide, x, y, w, headers: list[str], rows: list[list[str]], *, size=10, col_widths=None, row_h=Inches(0.36)
    ):
        shape = slide.shapes.add_table(len(rows) + 1, len(headers), x, y, w, row_h * (len(rows) + 1))
        t = shape.table
        if col_widths:
            for i, cw in enumerate(col_widths):
                t.columns[i].width = cw
        for j, hdr in enumerate(headers):
            c = t.cell(0, j)
            c.text = hdr
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size, r.font.bold, r.font.name = Pt(size), True, FONT
                    r.font.color.rgb = WHITE
        for i, row in enumerate(rows, start=1):
            for j, val in enumerate(row):
                c = t.cell(i, j)
                c.text = str(val)
                c.fill.solid()
                c.fill.fore_color.rgb = RGBColor(0xF5, 0xF5, 0xF7) if i % 2 else WHITE
                for p in c.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size, r.font.name = Pt(size), FONT
                        r.font.color.rgb = INK
        return shape

    def save(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path


def _join(items: list[str], n=3) -> str:
    return "; ".join(items[:n])


def render_pptx(
    *,
    request: ResearchRequest,
    analysis: AnalysisReport,
    products: list[ExtractedProduct],
    stats: PriceStats | None,
    images: list[ImageAsset],
    report_dir: Path,
    out_path: Path,
) -> Path:
    a = analysis
    d = Deck(a.title)

    # 1 cover
    s = d.prs.slides.add_slide(d.blank)
    d.n += 1
    bg = s.shapes.add_shape(1, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    d._text(
        s,
        Inches(0.9),
        Inches(1.6),
        Inches(11.5),
        Inches(0.4),
        "COMPETITOR & MARKET RESEARCH REPORT",
        size=12,
        color=RGBColor(0xBF, 0xD4, 0xE4),
    )
    d._text(s, Inches(0.9), Inches(2.1), Inches(11.5), Inches(2.0), a.title, size=36, bold=True, color=WHITE)
    sub = [f"หัวข้อ: {request.query}"]
    if request.target_competitors:
        sub.append(f"คู่แข่งเป้าหมาย: {', '.join(request.target_competitors)}")
    if request.our_brand:
        sub.append(f"แบรนด์ของเรา: {request.our_brand}")
    sub.append(datetime.now().strftime("%d/%m/%Y"))
    d._text(s, Inches(0.9), Inches(4.4), Inches(11.5), Inches(1.6), sub, size=14, color=RGBColor(0xDD, 0xE7, 0xEF))
    d._text(
        s,
        Inches(0.9),
        H - Inches(0.8),
        Inches(11.5),
        Inches(0.4),
        "AutoResearch Agent — Puwanath Baibua · สร้างอัตโนมัติด้วย LLM โปรดตรวจสอบตัวเลขก่อนนำไปใช้",
        size=9,
        color=RGBColor(0x9F, 0xB6, 0xC6),
    )

    # 2 executive summary
    s = d.slide("บทสรุปผู้บริหาร", "Executive Summary")
    d._text(s, Inches(0.6), Inches(1.5), Inches(7.4), Inches(4.5), a.executive_summary, size=14)
    kpis = [
        ("ราคามัธยฐานตลาด", f"฿{stats.median:,.0f}" if stats and stats.median else "—"),
        ("ช่วงราคา", f"฿{stats.min:,.0f} – ฿{stats.max:,.0f}" if stats and stats.min is not None else "—"),
        ("ช่วงราคาที่แนะนำ", a.recommended_price_range_thb),
        ("ความเชื่อมั่น", {"low": "ต่ำ", "medium": "ปานกลาง", "high": "สูง"}[a.confidence]),
    ]
    for i, (k, v) in enumerate(kpis):
        y = Inches(1.5) + i * Inches(1.15)
        box = s.shapes.add_shape(1, Inches(8.4), y, Inches(4.3), Inches(1.0))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xF5, 0xF5, 0xF7)
        box.line.fill.background()
        d._text(s, Inches(8.6), y + Inches(0.1), Inches(4), Inches(0.3), k, size=10, color=GREY)
        d._text(s, Inches(8.6), y + Inches(0.38), Inches(4), Inches(0.6), v, size=20, bold=True, color=NAVY)

    # 3 key findings
    s = d.slide("ข้อค้นพบสำคัญ", "Key Findings")
    d._bullets(s, Inches(0.6), Inches(1.5), Inches(12), Inches(5), a.key_findings, size=16)

    # 4 price chart (native, editable)
    rows = chart_brands(products, request.target_competitors, 12)
    if len(rows) >= 2:
        s = d.slide("เปรียบเทียบราคามัธยฐานตามแบรนด์ (THB)", "Price Benchmark")
        cd = CategoryChartData()
        cd.categories = [r["brand"] for r in rows]
        cd.add_series("ราคามัธยฐาน", [r["median_price"] for r in rows])
        cd.add_series("ราคาต่ำสุด", [r["min_price"] for r in rows])
        cd.add_series("ราคาสูงสุด", [r["max_price"] for r in rows])
        gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.4), Inches(8.2), Inches(5.5), cd)
        ch = gf.chart
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.plots[0].series[0].format.fill.solid()
        ch.plots[0].series[0].format.fill.fore_color.rgb = NAVY
        ch.plots[0].series[1].format.fill.solid()
        ch.plots[0].series[1].format.fill.fore_color.rgb = RGBColor(0xB9, 0xCF, 0xE0)
        ch.plots[0].series[2].format.fill.solid()
        ch.plots[0].series[2].format.fill.fore_color.rgb = ACCENT
        ch.category_axis.tick_labels.font.size = Pt(10)
        ch.value_axis.tick_labels.font.size = Pt(9)
        d._text(s, Inches(9.1), Inches(1.5), Inches(3.8), Inches(5), a.pricing_insight, size=12)

    # 5 competitors table
    s = d.slide("ภาพรวมคู่แข่ง", "Competitor Landscape")
    d._text(s, Inches(0.6), Inches(1.35), Inches(12), Inches(0.9), a.market_overview, size=11, color=GREY)
    d.table(
        s,
        Inches(0.6),
        Inches(2.35),
        Inches(12.1),
        ["แบรนด์", "ตำแหน่งราคา", "ช่วงราคา", "จุดเด่น", "จุดด้อย", "ช่องทาง"],
        [
            [
                c.brand_name,
                c.price_position,
                c.price_range_thb,
                _join(c.strengths, 2),
                _join(c.weaknesses, 2),
                ", ".join(c.channels[:3]),
            ]
            for c in a.competitors[:8]
        ],
        size=9,
        col_widths=[Inches(1.5), Inches(1.2), Inches(1.4), Inches(3.4), Inches(3.0), Inches(1.6)],
        row_h=Inches(0.5),
    )

    # 6 product gallery
    if images:
        s = d.slide("ตัวอย่างสินค้าของคู่แข่ง", "Products")
        per_row = 4
        cell_w = Inches(3.0)
        for i, im in enumerate(images[:8]):
            r, c = divmod(i, per_row)
            x = Inches(0.6) + c * cell_w
            y = Inches(1.5) + r * Inches(2.9)
            path = report_dir / im.local_file_path
            if path.exists():
                ratio = im.width_px / max(im.height_px, 1)
                h = Inches(1.9)
                w = min(Emu(int(h * ratio)), Inches(2.7))
                s.shapes.add_picture(str(path), x + (Inches(2.7) - w) / 2, y, width=w, height=h)
            d._text(
                s,
                x,
                y + Inches(1.95),
                Inches(2.7),
                Inches(0.7),
                [im.brand_name or "", im.product_name[:48]],
                size=9,
                align=PP_ALIGN.CENTER,
            )

    # 7 channels
    ca = a.channel_analysis
    s = d.slide("ช่องทางการจัดจำหน่าย", "Distribution Channels")
    d._text(
        s,
        Inches(0.6),
        Inches(1.35),
        Inches(12),
        Inches(0.8),
        (ca.summary if ca else a.channel_insight),
        size=11,
        color=GREY,
    )
    if ca and ca.stats:
        cd = CategoryChartData()
        cd.categories = [c.channel_name for c in ca.stats]
        cd.add_series("จำนวนรายการ", [c.listing_count for c in ca.stats])
        gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(2.2), Inches(5.6), Inches(4.5), cd)
        gf.chart.has_legend = False
        gf.chart.plots[0].series[0].format.fill.solid()
        gf.chart.plots[0].series[0].format.fill.fore_color.rgb = NAVY
        gf.chart.category_axis.tick_labels.font.size = Pt(9)
        d.table(
            s,
            Inches(6.5),
            Inches(2.2),
            Inches(6.2),
            ["ช่องทาง", "บทบาท", "เหมาะกับเรา", "คำแนะนำ"],
            [
                [
                    c.channel_name,
                    c.role[:60],
                    {"high": "สูง", "medium": "กลาง", "low": "ต่ำ"}[c.fit_for_us],
                    c.recommendation[:80],
                ]
                for c in ca.channels[:7]
            ],
            size=9,
            col_widths=[Inches(1.2), Inches(2.0), Inches(0.9), Inches(2.1)],
            row_h=Inches(0.55),
        )
    else:
        d._bullets(
            s,
            Inches(0.6),
            Inches(2.3),
            Inches(12),
            Inches(4),
            [f"ช่องทางที่แนะนำ: {', '.join(a.recommended_channels)}"],
            size=14,
        )
    if ca and ca.matrix:
        channels = [k for k in ca.matrix[0] if k not in ("brand", "total")]
        s = d.slide("Matrix แบรนด์ × ช่องทาง (จำนวนรายการที่พบ)", "Distribution Channels")
        d.table(
            s,
            Inches(0.6),
            Inches(1.5),
            Inches(12.1),
            ["แบรนด์", *channels, "รวม"],
            [[r["brand"], *[r[c] or "" for c in channels], r["total"]] for r in ca.matrix[:12]],
            size=10,
        )
        d._bullets(
            s,
            Inches(0.6),
            Inches(1.5) + Inches(0.36) * (min(len(ca.matrix), 12) + 1) + Inches(0.3),
            Inches(12),
            Inches(1.5),
            ca.channel_mix_recommendation[:4],
            size=11,
        )

    # 8 usage insights
    u = a.usage_insights
    if u:
        s = d.slide("พฤติกรรมและการใช้งานของผู้บริโภค", "Usage & Consumer Insights")
        d._text(s, Inches(0.6), Inches(1.35), Inches(12), Inches(0.8), u.summary, size=11, color=GREY)
        cols = [("การใช้งานหลัก", u.use_cases), ("เหตุผลในการซื้อ", u.purchase_drivers), ("ปัญหา / ข้อกังวล", u.pain_points)]
        for i, (title, items) in enumerate(cols):
            x = Inches(0.6) + i * Inches(4.1)
            d._text(s, x, Inches(2.2), Inches(3.9), Inches(0.4), title, size=13, bold=True, color=NAVY)
            d._bullets(s, x, Inches(2.6), Inches(3.9), Inches(2.4), items[:5], size=11)
        d._text(s, Inches(0.6), Inches(5.1), Inches(12), Inches(0.4), "กลุ่มเป้าหมาย", size=13, bold=True, color=NAVY)
        d._bullets(
            s,
            Inches(0.6),
            Inches(5.5),
            Inches(12),
            Inches(1.5),
            [
                f"{t.segment}: {_join(t.needs, 3)}" + (f" → {t.opportunity}" if t.opportunity else "")
                for t in u.target_segments[:4]
            ],
            size=11,
        )

    # 9 promotions + features
    pa, fc = a.promotion_analysis, a.feature_comparison
    if pa or fc:
        s = d.slide("โปรโมชั่นและคุณสมบัติสินค้า", "Promotions & Features")
        if pa:
            d._text(s, Inches(0.6), Inches(1.4), Inches(6), Inches(0.4), "กลยุทธ์โปรโมชั่น", size=13, bold=True, color=NAVY)
            d._bullets(
                s,
                Inches(0.6),
                Inches(1.8),
                Inches(6),
                Inches(5),
                (pa.brand_tactics[:5] + pa.recommendations[:3]),
                size=11,
            )
        if fc:
            d._text(
                s, Inches(6.9), Inches(1.4), Inches(6), Inches(0.4), "คุณสมบัติที่แข่งขันกัน", size=13, bold=True, color=NAVY
            )
            d.table(
                s,
                Inches(6.9),
                Inches(1.85),
                Inches(5.8),
                ["คุณสมบัติ", "แบรนด์ที่มี", "พื้นฐาน"],
                [
                    [f.feature[:40], ", ".join(f.brands_offering[:4]), "✓" if f.is_table_stakes else ""]
                    for f in fc.features[:8]
                ],
                size=9,
                col_widths=[Inches(2.4), Inches(2.7), Inches(0.7)],
                row_h=Inches(0.4),
            )

    # 10 SWOT
    s = d.slide("SWOT Analysis", "Strategy")
    quad = [
        ("Strengths", a.swot.strengths, RGBColor(0xE6, 0xF6, 0xEA)),
        ("Weaknesses", a.swot.weaknesses, RGBColor(0xFD, 0xEC, 0xEC)),
        ("Opportunities", a.swot.opportunities, RGBColor(0xE8, 0xF0, 0xFB)),
        ("Threats", a.swot.threats, RGBColor(0xFF, 0xF3, 0xE0)),
    ]
    for i, (title, items, color) in enumerate(quad):
        r, c = divmod(i, 2)
        x, y = Inches(0.6) + c * Inches(6.15), Inches(1.45) + r * Inches(2.75)
        box = s.shapes.add_shape(1, x, y, Inches(6.0), Inches(2.6))
        box.fill.solid()
        box.fill.fore_color.rgb = color
        box.line.fill.background()
        d._text(s, x + Inches(0.2), y + Inches(0.1), Inches(5.6), Inches(0.4), title, size=13, bold=True, color=NAVY)
        d._bullets(s, x + Inches(0.2), y + Inches(0.5), Inches(5.6), Inches(2.0), items[:4], size=11)

    # 11 recommendations
    s = d.slide("ข้อเสนอแนะเชิงกลยุทธ์และแผนปฏิบัติ", "Recommendations")
    d._text(s, Inches(0.6), Inches(1.4), Inches(6), Inches(0.4), "ข้อเสนอแนะ", size=13, bold=True, color=NAVY)
    d._bullets(s, Inches(0.6), Inches(1.8), Inches(6), Inches(5), a.recommendations[:6], size=12)
    d._text(s, Inches(6.9), Inches(1.4), Inches(6), Inches(0.4), "Action Plan", size=13, bold=True, color=NAVY)
    d._bullets(
        s,
        Inches(6.9),
        Inches(1.8),
        Inches(6),
        Inches(5),
        [f"{i + 1}. {x}" for i, x in enumerate(a.action_plan[:6])],
        size=12,
    )

    # 12 data gaps
    if a.data_gaps:
        s = d.slide("ข้อจำกัดของข้อมูล", "Appendix")
        d._bullets(s, Inches(0.6), Inches(1.5), Inches(12), Inches(5), a.data_gaps, size=13)

    return d.save(out_path)
