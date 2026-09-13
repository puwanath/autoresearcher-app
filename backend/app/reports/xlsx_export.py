"""Excel export (PRD §5.4): raw data split into sheets with hyperlinks back to sources and pre-computed stats."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.agents.stats import brand_price_table
from app.schemas.research import AnalysisReport, ExtractedProduct, PriceStats, ResearchRequest, ScrapedPage

HEADER_FILL = PatternFill("solid", fgColor="0B3D5C")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Tahoma", size=10)
BODY_FONT = Font(name="Tahoma", size=10)
LINK_FONT = Font(name="Tahoma", size=10, color="0B5C8A", underline="single")
THIN = Side(style="thin", color="D2D2D7")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
THB = '#,##0.00 "฿"'


def _sheet(
    wb: Workbook,
    title: str,
    headers: list[str],
    rows: list[list],
    *,
    widths: list[int] | None = None,
    money_cols: set[int] = frozenset(),
    link_col: int | None = None,
) -> Worksheet:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for i, c in enumerate(r, start=1):
            c.font, c.border = BODY_FONT, BORDER
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if i in money_cols and isinstance(c.value, (int, float)):
                c.number_format = THB
            if link_col == i and isinstance(c.value, str) and c.value.startswith("http"):
                c.hyperlink = c.value
                c.font = LINK_FONT
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def render_xlsx(
    *,
    request: ResearchRequest,
    analysis: AnalysisReport,
    products: list[ExtractedProduct],
    stats: PriceStats | None,
    pages: list[ScrapedPage],
    out_path: Path,
) -> Path:
    wb = Workbook()
    a = analysis

    # ---- Summary --------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = a.title
    ws["A1"].font = Font(bold=True, size=14, name="Tahoma", color="0B3D5C")
    info = [
        ("หัวข้อวิจัย", request.query),
        ("คู่แข่งเป้าหมาย", ", ".join(request.target_competitors) or "-"),
        ("แบรนด์ของเรา", request.our_brand or "-"),
        ("วันที่", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("ความเชื่อมั่น", a.confidence),
        ("จำนวนสินค้า", len(products)),
        ("ราคามัธยฐาน (คำนวณล่วงหน้า)", stats.median if stats else None),
        ("ราคาเฉลี่ย (สูตร live จากชีต Prices)", '=IFERROR(AVERAGE(Prices!E:E),"")'),
        ("ราคาต่ำสุด (สูตร)", '=IFERROR(MIN(Prices!E:E),"")'),
        ("ราคาสูงสุด (สูตร)", '=IFERROR(MAX(Prices!E:E),"")'),
        ("ช่วงราคาที่แนะนำ", a.recommended_price_range_thb),
        ("ช่องทางที่แนะนำ", ", ".join(a.recommended_channels)),
    ]
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(i, 1, k).font = Font(bold=True, name="Tahoma", size=10)
        c = ws.cell(i, 2, v)
        c.font = BODY_FONT
        if "ราคา" in k and "ช่วง" not in k:
            c.number_format = THB
    ws.cell(len(info) + 4, 1, "Executive Summary").font = Font(bold=True, name="Tahoma", size=11, color="0B3D5C")
    c = ws.cell(len(info) + 5, 1, a.executive_summary)
    c.font, c.alignment = BODY_FONT, Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=len(info) + 5, start_column=1, end_row=len(info) + 5, end_column=6)
    ws.row_dimensions[len(info) + 5].height = 120
    r0 = len(info) + 7
    ws.cell(r0, 1, "Key Findings").font = Font(bold=True, name="Tahoma", size=11, color="0B3D5C")
    for i, k in enumerate(a.key_findings, start=1):
        ws.cell(r0 + i, 1, f"{i}. {k}").font = BODY_FONT
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60

    # ---- Competitors ----------------------------------------------------
    _sheet(
        wb,
        "Competitors",
        ["แบรนด์", "ตำแหน่งราคา", "ช่วงราคา", "จุดเด่น", "จุดด้อย", "ช่องทาง", "Sentiment", "หมายเหตุ"],
        [
            [
                c.brand_name,
                c.price_position,
                c.price_range_thb,
                "\n".join(c.strengths),
                "\n".join(c.weaknesses),
                ", ".join(c.channels),
                c.sentiment,
                c.sentiment_note,
            ]
            for c in a.competitors
        ],
        widths=[18, 12, 16, 45, 45, 24, 10, 40],
    )

    # ---- Brand summary (pre-computed) -----------------------------------
    rows = brand_price_table(products)
    _sheet(
        wb,
        "BrandSummary",
        [
            "แบรนด์",
            "จำนวนสินค้า",
            "ราคาต่ำสุด",
            "ราคามัธยฐาน",
            "ราคาสูงสุด",
            "ช่องทาง",
            "โปรโมชั่น",
            "คะแนนรีวิวเฉลี่ย",
            "ส่วนแบ่งรายการ (%)",
        ],
        [
            [
                r["brand"],
                r["products"],
                r["min_price"],
                r["median_price"],
                r["max_price"],
                ", ".join(r["channels"]),
                ", ".join(r["promotions"]),
                r["avg_rating"],
                round(100 * r["products"] / max(len(products), 1), 1),
            ]
            for r in rows
        ],
        widths=[20, 12, 14, 14, 14, 30, 30, 14, 16],
        money_cols={3, 4, 5},
    )

    # ---- Products / Prices ---------------------------------------------
    _sheet(
        wb,
        "Prices",
        [
            "สินค้า",
            "แบรนด์",
            "Variant",
            "หมวด",
            "ราคา",
            "สกุลเงิน",
            "ราคาเดิม",
            "ส่วนลด %",
            "ราคาต่อหน่วย",
            "ช่องทาง",
            "สถานะสต็อก",
            "คะแนน",
            "จำนวนรีวิว",
            "โปรโมชั่น",
            "แหล่งข้อมูล",
        ],
        [
            [
                p.product_name,
                p.brand_name,
                p.variant,
                p.category,
                p.price.amount,
                p.price.currency,
                p.price.original_price,
                p.price.discount_percentage,
                p.price.unit_price,
                ", ".join(c.channel_name for c in p.sales_channels),
                ", ".join(c.availability for c in p.sales_channels if c.availability != "Unknown"),
                p.rating,
                p.review_count,
                "; ".join(f"{pr.promo_type}: {pr.promo_description}".strip(": ") for pr in p.promotions),
                p.data_source_url,
            ]
            for p in products
        ],
        widths=[44, 16, 16, 14, 12, 8, 12, 10, 12, 22, 14, 8, 10, 40, 50],
        money_cols={5, 7, 9},
        link_col=15,
    )

    # ---- Channels ------------------------------------------------------
    ca = a.channel_analysis
    if ca and ca.matrix:
        channels = [k for k in ca.matrix[0] if k not in ("brand", "total")]
        _sheet(
            wb,
            "ChannelMatrix",
            ["แบรนด์", *channels, "รวม"],
            [[r["brand"], *[r[c] for c in channels], r["total"]] for r in ca.matrix],
            widths=[20] + [14] * (len(channels) + 1),
        )
    if ca:
        _sheet(
            wb,
            "Channels",
            [
                "ช่องทาง",
                "รายการ",
                "แบรนด์",
                "ราคาต่ำสุด",
                "ราคามัธยฐาน",
                "ราคาสูงสุด",
                "คะแนนผู้ขายเฉลี่ย",
                "โปรโมชั่น",
                "บทบาท",
                "จุดแข็ง",
                "ข้อควรระวัง",
                "เหมาะกับเรา",
                "คำแนะนำ",
            ],
            [
                [
                    st.channel_name,
                    st.listing_count,
                    st.brand_count,
                    st.price_min,
                    st.price_median,
                    st.price_max,
                    st.avg_seller_rating,
                    st.promo_count,
                    *(
                        next(
                            (
                                [c.role, "\n".join(c.strengths), "\n".join(c.watchouts), c.fit_for_us, c.recommendation]
                                for c in ca.channels
                                if c.channel_name == st.channel_name
                            ),
                            ["", "", "", "", ""],
                        )
                    ),
                ]
                for st in ca.stats
            ],
            widths=[18, 10, 10, 13, 13, 13, 12, 10, 36, 36, 36, 12, 40],
            money_cols={4, 5, 6},
        )

    # ---- Promotions ----------------------------------------------------
    _sheet(
        wb,
        "Promotions",
        ["แบรนด์", "สินค้า", "ประเภทโปรโมชั่น", "รายละเอียด", "หมดอายุ", "แหล่งข้อมูล"],
        [
            [p.brand_name, p.product_name, pr.promo_type, pr.promo_description, pr.valid_until, p.data_source_url]
            for p in products
            for pr in p.promotions
        ],
        widths=[16, 40, 18, 50, 12, 50],
        link_col=6,
    )

    # ---- Insights ------------------------------------------------------
    u = a.usage_insights
    if u:
        rows = [["สรุป", u.summary]]
        rows += [["การใช้งาน", x] for x in u.use_cases] + [["โอกาสในการใช้", x] for x in u.usage_occasions]
        rows += [["เหตุผลในการซื้อ", x] for x in u.purchase_drivers] + [["Pain point", x] for x in u.pain_points]
        rows += [
            [
                f"กลุ่มเป้าหมาย: {t.segment}",
                f"ความต้องการ: {'; '.join(t.needs)} | แบรนด์ที่ตอบโจทย์: {', '.join(t.brands_serving)} | โอกาส: {t.opportunity}",
            ]
            for t in u.target_segments
        ]
        rows += [["หลักฐานจากรีวิว", x] for x in u.evidence]
        _sheet(wb, "UsageInsights", ["ประเภท", "รายละเอียด"], rows, widths=[24, 100])
    fc = a.feature_comparison
    if fc:
        _sheet(
            wb,
            "Features",
            ["คุณสมบัติ", "แบรนด์ที่มี", "เป็นพื้นฐานของตลาด"],
            [[f.feature, ", ".join(f.brands_offering), "ใช่" if f.is_table_stakes else ""] for f in fc.features],
            widths=[36, 60, 18],
        )

    # ---- Strategy / Sources ---------------------------------------------
    rows = (
        [["Strength", x] for x in a.swot.strengths]
        + [["Weakness", x] for x in a.swot.weaknesses]
        + [["Opportunity", x] for x in a.swot.opportunities]
        + [["Threat", x] for x in a.swot.threats]
    )
    rows += (
        [["Recommendation", x] for x in a.recommendations]
        + [["Action", x] for x in a.action_plan]
        + [["Data gap", x] for x in a.data_gaps]
    )
    _sheet(wb, "Strategy", ["ประเภท", "รายละเอียด"], rows, widths=[16, 110])
    _sheet(
        wb,
        "Sources",
        ["ชื่อหน้า", "URL", "สถานะ", "วิธีดึง", "ผลลัพธ์"],
        [[p.title, p.final_url, p.status_code, p.fetched_with, p.error or "ok"] for p in pages],
        widths=[50, 70, 8, 12, 20],
        link_col=2,
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    return out_path
