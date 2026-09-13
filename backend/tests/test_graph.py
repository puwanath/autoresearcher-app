import pytest

from app.agents.graph import Deps, run_research
from app.schemas.research import OutputFormat, ResearchRequest
from tests.conftest import FakeLLM, fake_fetch, fake_search


@pytest.mark.asyncio
async def test_full_loop_without_refine(settings):
    llm = FakeLLM(products_per_page=2)
    deps = Deps(llm=llm, settings=settings, search=fake_search, fetch=fake_fetch, fetch_images=False)
    req = ResearchRequest(query="ครีมกันแดด https://brand.example.com/sun", output_formats=[OutputFormat.markdown])

    result = await run_research(req, "t1", deps)

    assert result.plan and "https://brand.example.com/sun" in result.plan.direct_urls
    assert llm.refine_calls == 0  # 4 pages × 2 products ≥ threshold
    assert result.pages_fetched == settings.max_pages
    assert len(result.products) == 8 and all(p.data_source_url for p in result.products)
    assert {p.brand_name for p in result.products} == {"BrandA"}
    assert result.price_stats.sample_size == 8
    assert result.analysis.title == "Report"
    assert result.analysis.channel_analysis and result.analysis.channel_analysis.matrix[0]["brand"] == "BrandA"
    assert result.analysis.channel_analysis.stats[0].channel_name == "Shopee"
    assert result.analysis.usage_insights.target_segments[0].segment == "oily skin"
    assert not result.errors, result.errors
    md = result.report_markdown
    assert "## บทสรุปผู้บริหาร" in md and "| BrandA |" in md
    assert "Matrix แบรนด์ × ช่องทาง" in md and "ส่วนที่ 4: พฤติกรรมและการใช้งาน" in md
    assert "ส่วนที่ 5: ข้อเสนอแนะ" in md and "เปรียบเทียบคุณสมบัติ" in md
    assert len(result.artifacts) == 1 and result.artifacts[0].path.endswith(".md")
    assert any(e.split("] ")[1].startswith("analyze") for e in result.events)


@pytest.mark.asyncio
async def test_loop_refines_when_data_is_thin(settings):
    llm = FakeLLM(products_per_page=0)
    deps = Deps(llm=llm, settings=settings, search=fake_search, fetch=fake_fetch, fetch_images=False)
    fetched: list[list[str]] = []

    async def counting_fetch(urls, task_id):
        fetched.append(urls)
        return await fake_fetch(urls, task_id)

    deps.fetch = counting_fetch
    result = await run_research(ResearchRequest(query="niche thing"), "t2", deps)

    assert llm.refine_calls == settings.max_iterations  # bounded loop
    assert len(fetched) == settings.max_iterations + 1
    assert not set(fetched[0]) & set(fetched[1])  # refined pass only visits new URLs
    assert result.analysis is not None  # still produces a report on thin data
    assert result.price_stats.sample_size == 0


@pytest.mark.asyncio
async def test_refine_with_no_new_queries_skips_to_analysis(settings):
    llm = FakeLLM(products_per_page=0)
    orig = llm.chat_json

    async def same_queries(messages, schema, **kw):
        out = await orig(messages, schema, **kw)
        if schema.__name__ == "_RefinedQueries":
            out.search_queries = ["q1"]  # already used by the planner
        return out

    llm.chat_json = same_queries
    fetched = []

    async def counting_fetch(urls, task_id):
        fetched.append(urls)
        return await fake_fetch(urls, task_id)

    deps = Deps(llm=llm, settings=settings, search=fake_search, fetch=counting_fetch, fetch_images=False)
    result = await run_research(ResearchRequest(query="niche"), "t4", deps)
    assert llm.refine_calls == 1 and len(fetched) == 1 and result.analysis is not None


@pytest.mark.asyncio
async def test_pdf_artifact(settings):
    deps = Deps(llm=FakeLLM(), settings=settings, search=fake_search, fetch=fake_fetch, fetch_images=False)
    req = ResearchRequest(query="กาแฟดริป", output_formats=[OutputFormat.markdown, OutputFormat.pdf])
    result = await run_research(req, "t3", deps)
    formats = {a.format for a in result.artifacts}
    assert formats == {OutputFormat.markdown, OutputFormat.pdf}
    pdf = next(a for a in result.artifacts if a.format == OutputFormat.pdf)
    assert open(pdf.path, "rb").read(5) == b"%PDF-"


@pytest.mark.asyncio
async def test_pptx_and_xlsx_artifacts(settings):
    deps = Deps(llm=FakeLLM(), settings=settings, search=fake_search, fetch=fake_fetch, fetch_images=False)
    req = ResearchRequest(query="กาแฟดริป", output_formats=[OutputFormat.pptx, OutputFormat.xlsx])
    result = await run_research(req, "t5", deps)
    paths = {a.format: a.path for a in result.artifacts}
    assert set(paths) == {OutputFormat.markdown, OutputFormat.pptx, OutputFormat.xlsx}

    from pptx import Presentation

    prs = Presentation(paths[OutputFormat.pptx])
    titles = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
    assert len(prs.slides) >= 9
    assert any("บทสรุปผู้บริหาร" in t for t in titles) and any("SWOT" in t for t in titles)
    assert any("พฤติกรรมและการใช้งาน" in t for t in titles)
    assert any(sh.has_chart for s in prs.slides for sh in s.shapes)  # native, editable chart

    from openpyxl import load_workbook

    wb = load_workbook(paths[OutputFormat.xlsx])
    expected = {
        "Summary",
        "Competitors",
        "BrandSummary",
        "Prices",
        "ChannelMatrix",
        "Channels",
        "UsageInsights",
        "Strategy",
    }
    assert expected | {"Sources"} <= set(wb.sheetnames)
    prices = wb["Prices"]
    assert prices.max_row == 1 + len(result.products)
    assert prices.cell(2, 15).hyperlink is not None  # source link
    assert wb["Summary"]["B10"].value.startswith("=IFERROR(AVERAGE(")
